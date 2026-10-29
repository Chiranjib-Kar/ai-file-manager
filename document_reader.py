import os
import re
import math
from typing import Dict, Any

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".log",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".html", ".css",
    ".xml", ".yaml", ".yml", ".sql", ".sh", ".bat", ".ini",
    ".conf", ".env", ".rst"
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp", ".ico"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a", ".flac"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm"}
ARCHIVE_EXTENSIONS = {".zip", ".tar", ".gz", ".rar", ".7z"}

def categorize_file_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext in {".pdf", ".doc", ".docx", ".txt", ".rtf", ".odt", ".epub", ".pages", ".md"}:
        return "Document"
    if ext in IMAGE_EXTENSIONS:
        return "Image"
    if ext in {".py", ".js", ".jsx", ".ts", ".tsx", ".html", ".css", ".json", ".xml", ".yaml", ".yml", ".sql", ".sh", ".bat"}:
        return "Code & Data"
    if ext in {".csv", ".xls", ".xlsx", ".tsv"}:
        return "Spreadsheet"
    if ext in ARCHIVE_EXTENSIONS:
        return "Archive"
    if ext in AUDIO_EXTENSIONS or ext in VIDEO_EXTENSIONS:
        return "Media"
    return "Other"

def extract_document_text(file_path: str, max_chars: int = 15000) -> Dict[str, Any]:
    """
    Extracts text, calculates metrics, and inspects structure of document.
    """
    ext = os.path.splitext(file_path)[1].lower()
    extracted_text = ""
    metadata = {}

    if not os.path.exists(file_path):
        return {
            "text": "",
            "preview": "File not found.",
            "word_count": 0,
            "reading_time_minutes": 0,
            "metadata": {}
        }

    # 1. PDF Documents
    if ext == ".pdf":
        if pypdf:
            try:
                reader = pypdf.PdfReader(file_path)
                pages_text = []
                num_pages = len(reader.pages)
                metadata["pages"] = num_pages
                if reader.metadata:
                    if reader.metadata.title:
                        metadata["pdf_title"] = str(reader.metadata.title)
                    if reader.metadata.author:
                        metadata["pdf_author"] = str(reader.metadata.author)

                for idx, page in enumerate(reader.pages[:20]): # read up to first 20 pages
                    t = page.extract_text()
                    if t:
                        pages_text.append(t)
                extracted_text = "\n\n".join(pages_text)
            except Exception as e:
                extracted_text = f"[PDF Parsing Notice: {str(e)}]"
        else:
            extracted_text = "[PDF parser not available]"

    # 2. DOCX Documents
    elif ext in {".docx", ".doc"}:
        if docx and ext == ".docx":
            try:
                doc = docx.Document(file_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also read tables
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs.append(row_text)
                extracted_text = "\n\n".join(paragraphs)
                metadata["paragraphs"] = len(doc.paragraphs)
            except Exception as e:
                extracted_text = f"[DOCX Parsing Notice: {str(e)}]"
        else:
            extracted_text = "[Word docx parser fallback or binary format]"

    # 3. Plain Text / Markdown / Code / CSV / JSON
    elif ext in TEXT_EXTENSIONS:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                extracted_text = f.read(max_chars)
        except Exception as e:
            try:
                with open(file_path, "r", encoding="latin-1", errors="replace") as f:
                    extracted_text = f.read(max_chars)
            except Exception as e2:
                extracted_text = f"[Text Read Error: {str(e2)}]"

    # 4. Images / Other
    elif ext in IMAGE_EXTENSIONS:
        file_size = os.path.getsize(file_path)
        extracted_text = f"[Image Asset: {os.path.basename(file_path)} - {file_size} bytes. Visual content file.]"
        metadata["is_image"] = True
    else:
        # Fallback inspection
        try:
            with open(file_path, "rb") as f:
                header = f.read(1024)
                # Try simple ascii decode
                ascii_parts = re.findall(r'[ -~]{4,}', header.decode('latin-1', errors='ignore'))
                if ascii_parts:
                    extracted_text = " ".join(ascii_parts[:30])
                else:
                    extracted_text = f"[Binary/Data File: {os.path.basename(file_path)}]"
        except Exception:
            extracted_text = f"[File: {os.path.basename(file_path)}]"

    # Metrics
    cleaned_text = re.sub(r'\s+', ' ', extracted_text).strip()
    words = cleaned_text.split() if cleaned_text else []
    word_count = len(words)
    reading_time = max(1, math.ceil(word_count / 200)) if word_count > 0 else 0

    preview = cleaned_text[:350] + ("..." if len(cleaned_text) > 350 else "")

    return {
        "text": extracted_text[:max_chars],
        "preview": preview,
        "word_count": word_count,
        "reading_time_minutes": reading_time,
        "metadata": metadata
    }
